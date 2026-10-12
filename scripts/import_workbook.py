from __future__ import annotations

import argparse
import csv
import hashlib
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


FIELD_ORDER = [
    "priority","subject","status","charging_agency","us_attorneys_office","court",
    "case_number","charges_statutes","factual_allegations_or_admissions",
    "victims_or_public_harm","trigger_or_referral",
    "cooperators_whistleblowers_competitors_political","warrants_subpoenas_raids",
    "key_dates","potential_sentence","primary_source","reporting_source",
    "jrf_relevance","proof_posture","healthcare","public_professional_email",
    "public_counsel_or_contact","counsel_email","public_phone","outreach_status",
    "tags","first_seen","last_updated","fingerprint",
]

COLUMN_MAP = {
    "Priority": "priority",
    "Defendant / Subject": "subject",
    "Status": "status",
    "Charging / Investigative Agencies": "charging_agency",
    "U.S. Attorney / DOJ Component": "us_attorneys_office",
    "Court / District": "court",
    "Case No.": "case_number",
    "Charges / Statutes": "charges_statutes",
    "Allegations / Admitted Facts": "factual_allegations_or_admissions",
    "Alleged Victims / Public Harm": "victims_or_public_harm",
    "Investigative Trigger / Referral": "trigger_or_referral",
    "Cooperators / Whistleblowers / Competitors / Political Actors": "cooperators_whistleblowers_competitors_political",
    "Confirmed Search Warrants / Subpoenas / Raids": "warrants_subpoenas_raids",
    "Key Dates": "key_dates",
    "Potential Sentence / Statutory Maximum": "potential_sentence",
    "Primary Source": "primary_source",
    "Credible Reporting": "reporting_source",
    "JRF / P2P Relevance": "jrf_relevance",
    "Public Professional Email": "public_professional_email",
    "Public Counsel / Contact": "public_counsel_or_contact",
    "Counsel Email": "counsel_email",
    "Public Phone": "public_phone",
    "Outreach Status": "outreach_status",
    "Last Updated": "last_updated",
}

HEALTHCARE_TERMS = (
    "healthcare","health care","medicare","medicaid","physician","doctor",
    "hospital","laboratory","lab ","telehealth","telemedicine","pharmacy",
    "hhs-oig"," dme","durable medical","false claims act","qui tam",
)


def clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def proof_posture(status: str, proof_text: str) -> str:
    s = status.lower()
    p = proof_text.lower()
    if "guilty plea" in s or "pleaded guilty" in s or "admitted" in p:
        return "admitted"
    if "convicted" in s or "sentenced" in s or "conviction" in s or "proven at trial" in p:
        return "proven"
    if "mixed" in p:
        return "mixed"
    return "allegation"


def healthcare_flag(values: list[str]) -> str:
    text = " ".join(values).lower()
    return "true" if any(term in text for term in HEALTHCARE_TERMS) else "false"


def fp(row: dict[str, str]) -> str:
    stable = "|".join([
        row["subject"].lower(),
        row["case_number"].lower(),
        row["primary_source"].lower(),
        row["charges_statutes"].lower(),
    ])
    return hashlib.sha256(stable.encode("utf-8")).hexdigest()[:20]


def main() -> None:
    parser = argparse.ArgumentParser(description="Import JRF Master Registry workbook into data/registry.csv")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--sheet", default="Master Registry")
    parser.add_argument("--output", type=Path, default=Path("data/registry.csv"))
    args = parser.parse_args()

    wb = load_workbook(args.workbook, read_only=True, data_only=True)
    if args.sheet not in wb.sheetnames:
        raise SystemExit(f"Sheet {args.sheet!r} not found. Available: {wb.sheetnames}")
    ws = wb[args.sheet]

    iterator = ws.iter_rows(values_only=True)
    headers = [clean(v) for v in next(iterator)]
    index = {name: i for i, name in enumerate(headers)}

    missing = [name for name in COLUMN_MAP if name not in index and name not in {"Public Professional Email","Public Counsel / Contact","Counsel Email","Public Phone"}]
    if missing:
        raise SystemExit(f"Missing expected columns: {missing}")

    rows: list[dict[str, str]] = []
    for raw in iterator:
        get = lambda col: clean(raw[index[col]]) if col in index and index[col] < len(raw) else ""
        subject = get("Defendant / Subject")
        if not subject:
            continue

        row = {field: "" for field in FIELD_ORDER}
        for source_col, target_col in COLUMN_MAP.items():
            row[target_col] = get(source_col)

        proof_text = get("Proof Posture")
        row["proof_posture"] = proof_posture(row["status"], proof_text)
        row["healthcare"] = healthcare_flag([clean(v) for v in raw])
        row["tags"] = ""
        row["priority"] = row["priority"] or "MEDIUM"
        row["outreach_status"] = row["outreach_status"] or "RESEARCH"
        row["last_updated"] = row["last_updated"] or ""
        row["first_seen"] = row["last_updated"]
        row["fingerprint"] = fp(row)
        rows.append(row)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=FIELD_ORDER)
        writer.writeheader()
        writer.writerows(rows)

    print(f"Imported {len(rows)} registry entries to {args.output}")


if __name__ == "__main__":
    main()
